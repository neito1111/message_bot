from aiogram import Router, F
from aiogram.types import CallbackQuery, FSInputFile
from sqlalchemy import select
from bot.db import DMProfile, TGAccount, Chat, Event, EventType, User, async_session_maker
import logging
import tempfile
import os
import asyncio
import json
from datetime import datetime, timedelta, timezone
from collections import defaultdict

logger = logging.getLogger(__name__)
router = Router()
UKRAINE_FIXED_TZ = timezone(timedelta(hours=2))

LOADING_EMOJIS = [
    "\U0001F311",
    "\U0001F312",
    "\U0001F313",
    "\U0001F314",
    "\U0001F315",
    "\U0001F316",
    "\U0001F317",
    "\U0001F318",
]


@router.callback_query(F.data == "admin_export_stats")
async def export_stats_handler(callback: CallbackQuery):
    loading_message = await callback.message.answer("\U0001F311 \u0413\u0435\u043D\u0435\u0440\u0430\u0446\u0438\u044F \u043E\u0442\u0447\u0435\u0442\u0430...")
    await callback.answer()

    temp_file = None
    try:
        temp_file = await generate_stats_excel(loading_message)
        filename = f"stats_{asyncio.get_event_loop().time()}.xlsx"
        input_file = FSInputFile(temp_file, filename=filename)
        await callback.message.answer_document(
            document=input_file,
            caption="\u0421\u0442\u0430\u0442\u0438\u0441\u0442\u0438\u043A\u0430 \u0432\u044B\u0433\u0440\u0443\u0436\u0435\u043D\u0430",
        )
    except Exception as e:
        logger.error("Error exporting stats: %s", e)
        try:
            await loading_message.edit_text(f"\u041E\u0448\u0438\u0431\u043A\u0430 \u043F\u0440\u0438 \u0433\u0435\u043D\u0435\u0440\u0430\u0446\u0438\u0438 \u0444\u0430\u0439\u043B\u0430: {str(e)[:200]}")
        except Exception:
            await callback.message.answer(f"\u041E\u0448\u0438\u0431\u043A\u0430: {str(e)[:200]}")
    finally:
        if temp_file and os.path.exists(temp_file):
            try:
                os.unlink(temp_file)
            except Exception:
                pass
        try:
            await loading_message.delete()
        except Exception:
            pass


def parse_event_data(raw: str | None) -> dict:
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except Exception:
        return {"raw": raw}


def format_ukraine_time(dt) -> str:
    if not dt:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(UKRAINE_FIXED_TZ).strftime("%Y-%m-%d %H:%M:%S")


def to_ukraine_datetime(dt):
    if not dt:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(UKRAINE_FIXED_TZ)


def sanitize_sheet_title(title: str, fallback: str) -> str:
    cleaned = "".join("_" if ch in '[]:*?/\\' else ch for ch in (title or "")).strip()
    cleaned = cleaned[:31].strip()
    return cleaned or fallback


def append_rows_with_day_separators(ws, headers: list[str], rows: list[dict]):
    ws.append(headers)
    current_day = None

    for row in rows:
        row_day = row.pop("_day", None)
        if row_day and row_day != current_day:
            separator_row_index = ws.max_row + 1
            ws.append([str(row_day)] + [""] * (len(headers) - 1))
            ws.merge_cells(start_row=separator_row_index, start_column=1, end_row=separator_row_index, end_column=len(headers))
            separator_cell = ws.cell(row=separator_row_index, column=1)
            separator_cell.fill = separator_cell.fill.copy(fill_type="solid", fgColor="000000")
            separator_cell.font = separator_cell.font.copy(bold=True, color="FFFFFF")
            separator_cell.alignment = separator_cell.alignment.copy(horizontal="center")
            current_day = row_day

        ws.append([row.get(header, "") for header in headers])


async def generate_stats_excel(loading_message=None) -> str:
    try:
        from openpyxl import Workbook
    except ImportError:
        raise ImportError("openpyxl \u043D\u0435 \u0443\u0441\u0442\u0430\u043D\u043E\u0432\u043B\u0435\u043D. \u0423\u0441\u0442\u0430\u043D\u043E\u0432\u0438\u0442\u0435: pip install openpyxl")

    from openpyxl.styles import PatternFill, Font, Alignment

    wb = Workbook()
    wb.remove(wb.active)

    async with async_session_maker() as session:
        event_stmt = (
            select(Event, Chat, TGAccount, User, DMProfile)
            .join(Chat, Event.chat_id == Chat.id)
            .outerjoin(TGAccount, Chat.tg_account_id == TGAccount.id)
            .outerjoin(User, TGAccount.user_id == User.id)
            .outerjoin(DMProfile, DMProfile.user_id == User.id)
            .order_by(Event.created_at.desc(), Event.id.desc())
        )
        event_result = await session.execute(event_stmt)
        event_rows = event_result.all()

        summary_stmt = (
            select(Chat, TGAccount, User, DMProfile, Event)
            .join(TGAccount, Chat.tg_account_id == TGAccount.id)
            .join(User, TGAccount.user_id == User.id)
            .outerjoin(DMProfile, DMProfile.user_id == User.id)
            .outerjoin(Event, Event.chat_id == Chat.id)
            .order_by(Chat.created_at.desc(), Chat.id.desc(), Event.created_at.asc(), Event.id.asc())
        )
        summary_result = await session.execute(summary_stmt)
        summary_rows = summary_result.all()

        summary_by_chat: dict[int, dict] = {}
        for chat, tg_account, manager_user, dm_profile, event in summary_rows:
            item = summary_by_chat.get(chat.id)
            if item is None:
                manager_username = f"@{manager_user.username}" if manager_user and manager_user.username else "N/A"
                lead_username = chat.user_username or "N/A"
                if lead_username != "N/A" and not str(lead_username).startswith("@"):
                    lead_username = f"@{lead_username}"
                item = {
                    "Chat ID": chat.chat_id,
                    "Lead TG ID": chat.user_tg_id or "N/A",
                    "Lead Username": lead_username,
                    "Lead Name": chat.user_first_name or "N/A",
                    "Manager TG ID": manager_user.tg_id if manager_user else "N/A",
                    "Manager Username": manager_username,
                    "Manager Tag": dm_profile.tag if dm_profile and dm_profile.tag else "N/A",
                    "Manager Phone": tg_account.phone if tg_account and tg_account.phone else "N/A",
                    "Chat Title": chat.chat_title or "N/A",
                    "Messages Total": chat.messages_count or 0,
                    "REGA Time": "",
                    "REGA": chat.first_message_text or "",
                    "DEP Time": "",
                    "DEP": "",
                    "DO_DEP Time": "",
                    "DO_DEP": "",
                }
                summary_by_chat[chat.id] = item

            if event and event.event_type:
                event_data = parse_event_data(event.event_data)
                event_time = format_ukraine_time(event.created_at)
                if event.event_type == EventType.REGA:
                    item["REGA Time"] = event_time
                    item["REGA"] = event_data.get("message") or chat.first_message_text or ""
                elif event.event_type == EventType.DEP:
                    item["DEP Time"] = event_time
                    item["DEP"] = event_data.get("message_text") or ""
                elif event.event_type == EventType.DO_DEP:
                    item["DO_DEP Time"] = event_time
                    item["DO_DEP"] = event_data.get("trigger_message") or ""

        summary_data = []
        for item in summary_by_chat.values():
            if not any((item["REGA Time"], item["DEP Time"], item["DO_DEP Time"])):
                continue
            row = dict(item)
            summary_data.append(row)

        def effective_dt(row: dict):
            for key in ("DO_DEP Time", "DEP Time", "REGA Time"):
                value = row.get(key)
                if value:
                    return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
            return datetime.min

        summary_data.sort(key=effective_dt)

        summary_headers = [
            "Lead TG ID",
            "Lead Username",
            "Lead Name",
            "Manager TG ID",
            "Manager Username",
            "Manager Tag",
            "Chat Title",
            "Messages Total",
            "REGA Time",
            "REGA",
            "DEP Time",
            "DEP",
            "DO_DEP Time",
            "DO_DEP",
        ]
        for row in summary_data:
            row["_day"] = effective_dt(row).date() if effective_dt(row) != datetime.min else None

        manager_sheets: dict[str, list[dict]] = defaultdict(list)
        for row in summary_data:
            manager_key = row["Manager Tag"] if row["Manager Tag"] != "N/A" else row["Manager Username"]
            manager_sheets[manager_key].append(dict(row))

        sheets = [("Summary", summary_data)]
        for manager_key, manager_rows in sorted(manager_sheets.items(), key=lambda item: item[0]):
            title = sanitize_sheet_title(manager_key, "DM")
            sheets.append((title, manager_rows))

        header_fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        header_font = Font(bold=True)
        separator_fill = PatternFill(fill_type="solid", fgColor="000000")
        separator_font = Font(bold=True, color="FFFFFF")
        separator_alignment = Alignment(horizontal="center")
        frame = 0
        for sheet_name, data in sheets:
            if loading_message:
                emoji = LOADING_EMOJIS[frame % len(LOADING_EMOJIS)]
                try:
                    await loading_message.edit_text(f"{emoji} \u0413\u0435\u043D\u0435\u0440\u0430\u0446\u0438\u044F \u043B\u0438\u0441\u0442\u0430: {sheet_name}...")
                except Exception:
                    pass
                frame += 1
                await asyncio.sleep(0.2)

            ws = wb.create_sheet(title=sheet_name)
            if data:
                headers = summary_headers
                append_rows_with_day_separators(ws, headers, [dict(row) for row in data])
            else:
                ws.append(["No data"])

            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                if len(row) == 1 and row[0].value:
                    row[0].fill = separator_fill
                    row[0].font = separator_font
                    row[0].alignment = separator_alignment

        if loading_message:
            try:
                await loading_message.edit_text("\u2705 \u0413\u043E\u0442\u043E\u0432\u043E! \u041E\u0442\u043F\u0440\u0430\u0432\u043B\u044F\u044E \u0444\u0430\u0439\u043B...")
            except Exception:
                pass
            await asyncio.sleep(0.5)

        temp_fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(temp_fd)
        wb.save(temp_path)
        return temp_path
